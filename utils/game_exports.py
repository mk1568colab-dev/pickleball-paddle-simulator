"""Administrator-only, read-only classroom snapshots and analysis exports."""

from contextlib import closing, contextmanager
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import sqlite3
import tempfile
import time
from zipfile import ZIP_DEFLATED, ZipFile

from utils.classrooms import AUTH_CLASSROOM_KEY, DEFAULT_CLASSROOM, _session, active_classroom_id
from utils.database import get_base_database_path, get_database_path


BACKUP_FORMAT = "kiki-classroom-backup"
BACKUP_VERSION = 1
MAX_BACKUP_BYTES = 256 * 1024 * 1024
SHEETS = {
    "decision_audit_events": ("Decision Audit Events", "Successful saved revisions since logging was enabled; includes unchanged resubmissions"),
    "decision_audit_changes": ("Decision Audit Changes", "Field-level before/after values for each saved revision"),
    "decision_audit_context": ("Audit Coverage", "Logging start and current game ID; no historical backfill"),
    "round_results": ("Team Results", "All saved rounds"),
    "product_round_results": ("Product Results", "All saved rounds"),
    "team_decisions": ("Firm Decisions", "Last saved submission per team and round"),
    "product_decisions": ("Product Decisions", "Last saved submission per product and round"),
    "product_forecasts": ("Demand Forecasts", "All saved rounds"),
    "forecast_accuracy_results": ("Forecast Accuracy", "All saved rounds"),
    "market_reports": ("Market Reports", "All stored rounds, including unrevealed actual demand"),
    "demand_commitments": ("Demand Commitments", "Outlooks, committed demand and lock timestamps"),
    "classroom_rounds": ("Round Status", "Submission windows and instructor notes"),
    "material_rules": ("Material Rules", "Stored round rules; latest prior rule applies when absent"),
    "delivery_policy": ("Delivery Policy", "Opt-in rules: v1 per-supplier fee; v2 base fee plus additional-supplier fee. Replay seed is included only in the restricted backup"),
    "team_states": ("Current Team State", "Current snapshot, not a state-change history"),
    "product_lines": ("Current Products", "Current snapshot, not a portfolio-change history"),
    "product_development_projects": ("Current Development", "Current snapshot; overwritten project edits are not retained"),
    "team_archetypes": ("Starting Profiles", "Current operating baseline definitions"),
    "users": ("Account Roster", "Current accounts; password hashes excluded from Excel"),
}
ROSTER_FIELDS = {"user_id", "username", "role", "team_name", "is_active", "created_at", "updated_at"}


def _quoted(identifier):
    return '"' + identifier.replace('"', '""') + '"'


def _table_names(connection):
    return [row[0] for row in connection.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]


def _counts(connection):
    return {name: connection.execute(f"SELECT COUNT(*) FROM {_quoted(name)}").fetchone()[0]
            for name in _table_names(connection)}


def _classroom_metadata():
    classroom_id = active_classroom_id()
    base = get_base_database_path()
    registry = base.parent / (base.stem + "_classrooms.db")
    if not registry.is_file():
        if classroom_id != DEFAULT_CLASSROOM:
            raise PermissionError("Classroom is unavailable.")
        return classroom_id, "Class 1"
    with closing(sqlite3.connect(registry.resolve().as_uri() + "?mode=ro", uri=True)) as connection:
        row = connection.execute("SELECT name FROM classrooms WHERE classroom_id=? AND ready=1",
                                 (classroom_id,)).fetchone()
    if row is None:
        raise PermissionError("Classroom is unavailable.")
    return classroom_id, row[0]


@contextmanager
def _authorized_source():
    # Derive the scope from authentication, never from a supplied classroom/path.
    session = _session()
    username = session.get("authenticated_username") if session is not None else None
    if not username or session.get(AUTH_CLASSROOM_KEY, DEFAULT_CLASSROOM) != active_classroom_id():
        raise PermissionError("Sign in as this classroom's administrator.")
    path = get_database_path().resolve()
    with closing(sqlite3.connect(path.as_uri() + "?mode=ro", uri=True, timeout=15)) as connection:
        connection.execute("PRAGMA query_only=ON")
        connection.execute("BEGIN")
        user = connection.execute(
            "SELECT role, is_active FROM users WHERE username=?", (username,)).fetchone()
        if user is None or user[0] != "admin" or not user[1]:
            raise PermissionError("An active administrator in this classroom is required.")
        yield connection


def export_admin_user():
    """Authorize the export page without bootstrap migrations or seed updates."""
    from models.schemas import AppUser
    with _authorized_source() as connection:
        connection.row_factory = sqlite3.Row
        row = connection.execute(
            "SELECT user_id, username, role, team_name, is_active FROM users WHERE username=?",
            (_session()["authenticated_username"],)).fetchone()
        return AppUser.from_dict(dict(row))


def _plain_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        # Student-authored names and event text must never become Excel formulas.
        cell.data_type = "s"
    return cell


def _style_sheet(sheet, widths=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    sheet.freeze_panes = "C2" if sheet.max_column > 2 else "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 46
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="175C54")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, cell in enumerate(sheet[1], 1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            (widths or {}).get(index, min(35, max(18, len(str(cell.value)) * .9))))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="263238")
            cell.alignment = Alignment(vertical="top")
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00########"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"


def _workbook_bytes(connection, metadata):
    # The deployed Python app already uses openpyxl for its Excel importer.
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Export Overview"
    details = [
        ["Item", "Value"],
        ["Classroom", metadata["classroom_name"]],
        ["Classroom ID", metadata["classroom_id"]],
        ["Captured at (UTC)", metadata["captured_at_utc"]],
        ["Latest completed round", metadata["latest_completed_round"]],
        ["Scope", "One classroom. All stored rows, not only the latest round."],
        ["Privacy", "Administrator only. Contains account names and unrevealed market demand. No password hashes."],
        ["History limits", "Unsaved drafts and edits before logging are unavailable. Development and portfolio sheets are current snapshots; later saved revisions are in the audit sheets."],
        ["Values", "Saved values, not recalculations. Field names and units retain their database meaning."],
        ["Embedded records", "JSON columns preserve saved structures. Open Material Orders expands current outstanding orders."],
        ["Long text", "Text over 32,767 characters is split in Text Chunks. Unsupported control characters are escaped as \\uXXXX."],
        ["Restoration", "This Excel is for analysis. Use the separate backup ZIP for recovery."],
        ["Decision audit", "Events and field changes start when logging is enabled. Student, admin and automatic sources are distinct. Unsubmitted edits are not recorded."],
    ]
    for r, values in enumerate(details, 1):
        for c, value in enumerate(values, 1):
            _plain_cell(overview, r, c, value)
    _style_sheet(overview, {1: 27, 2: 110})
    for row in overview.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
        overview.row_dimensions[row[0].row].height = 34

    index = workbook.create_sheet("Table Index")
    index.append(["Worksheet", "Source table", "Rows", "Coverage"])
    chunks = []
    available = _table_names(connection)
    ordered = [name for name in SHEETS if name in available] + [name for name in available if name not in SHEETS]
    for table in ordered:
        title, scope = SHEETS.get(table, (table[:31], "All stored rows"))
        columns = [row[1] for row in connection.execute(f"PRAGMA table_info({_quoted(table)})")]
        if table == "users":
            columns = [name for name in columns if name in ROSTER_FIELDS]
        elif table == "delivery_policy":
            columns = [name for name in columns if name != 'seed']
        else:
            columns = [name for name in columns if not re.search("password|secret|token|credential", name, re.I)]
        if not columns:
            continue
        if len(columns) > 16384 or metadata["table_counts"][table] >= 1048576:
            raise ValueError("Saved data exceeds Excel worksheet limits. Use the backup ZIP for this classroom.")
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        order = [name for name in ("event_sequence", "change_id", "round_number", "team_name", "slot_name", "project_slot_name", "user_id") if name in columns]
        query = f"SELECT {', '.join(_quoted(name) for name in columns)} FROM {_quoted(table)}"
        if order:
            query += " ORDER BY " + ", ".join(_quoted(name) for name in order)
        for r, values in enumerate(connection.execute(query), 2):
            for c, value in enumerate(values, 1):
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m.group()):04x}", value)
                    if len(value) > 32767:
                        chunks.extend([table, r, columns[c-1], part // 30000 + 1, value[part:part+30000]]
                                      for part in range(0, len(value), 30000))
                        value = f"See Text Chunks: {table}, Excel row {r}, {columns[c-1]}"
                elif isinstance(value, int) and abs(value) >= 10 ** 15:
                    value = str(value)
                _plain_cell(sheet, r, c, value)
        _style_sheet(sheet)
        index.append([sheet.title, table, metadata["table_counts"][table], scope])

    if "team_states" in available:
        orders = workbook.create_sheet("Open Material Orders")
        rows = []
        keys = []
        for team, raw in connection.execute("SELECT team_name, open_material_orders_json FROM team_states ORDER BY team_name"):
            for position, order in enumerate(json.loads(raw or "[]"), 1):
                keys.extend(k for k in order if k not in keys)
                rows.append((team, position, order))
        orders.append(["team_name", "order_index"] + keys)
        for r, (team, position, order) in enumerate(rows, 2):
            for c, value in enumerate([team, position] + [order.get(k) for k in keys], 1):
                _plain_cell(orders, r, c, value)
        _style_sheet(orders)
        index.append([orders.title, "team_states.open_material_orders_json", len(rows), "Current outstanding orders only"])
    if chunks:
        sheet = workbook.create_sheet("Text Chunks")
        sheet.append(["source_table", "source_excel_row", "source_column", "part", "text"])
        for r, values in enumerate(chunks, 2):
            for c, value in enumerate(values, 1):
                _plain_cell(sheet, r, c, value)
        _style_sheet(sheet)
    _style_sheet(index, {1: 28, 2: 48, 3: 12, 4: 94})
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


RESTORE_README = """Kiki Pickleball - single-classroom recovery backup

PRIVATE: contains account password hashes, student identifiers, and unrevealed demand.
This ZIP is not encrypted. Store it securely; never share it with students or commit it.

Contents: classroom.sqlite3 (complete consistent SQLite snapshot), manifest.json,
and this file. Other classrooms and the server-wide registry are NOT included.
Code, assets, environment secrets and hosting settings are NOT included.
Use the matching application source/version; the database alone is not the app.

Safe recovery rehearsal, from the application source folder:
  python scripts/restore_classroom_backup.py BACKUP.zip NEW_FOLDER/recovered.db
The utility checks the checksum and SQLite integrity, and refuses existing targets.
It never registers, replaces, resets or migrates an existing classroom.

To inspect that new standalone copy in PowerShell:
  $env:SIMULATOR_DB_PATH = (Resolve-Path 'NEW_FOLDER/recovered.db').Path
  $env:SIMULATOR_ENV = 'dev'
  $env:SIMULATOR_ENABLE_DEMO_ACCOUNTS = 'false'
  python -m streamlit run app.py --server.address=127.0.0.1 --server.port=8555
Use an unused port. Select Class 1 and sign in with the backed-up admin account.
As a standalone copy its selector says Class 1; the original name/ID are in the manifest.

Restoring into a running server is deliberately not automated. Stop writers, retain
a new backup of each affected class and its registry, and verify the target class
before a separate, supervised replacement. Never overwrite a running SQLite file.
"""


def prepare_classroom_downloads():
    """Return session-local downloads made from one authorized consistent snapshot."""
    with _authorized_source() as source:
        classroom_id, name = _classroom_metadata()
        with tempfile.TemporaryDirectory(prefix="kiki-export-") as directory:
            path = Path(directory) / "classroom.sqlite3"
            deadline = time.monotonic() + 30

            def progress(status, remaining, total):
                if time.monotonic() > deadline:
                    raise TimeoutError("Snapshot timed out. Try again after the current round finishes.")

            with closing(sqlite3.connect(path)) as snapshot:
                source.backup(snapshot, pages=256, progress=progress, sleep=.05)
                # Standalone backup must not depend on a WAL or SHM sidecar.
                snapshot.execute("PRAGMA journal_mode=DELETE")
                if snapshot.execute("PRAGMA integrity_check").fetchall() != [("ok",)]:
                    raise ValueError("Database integrity check failed; no download was prepared.")
                unfinished = snapshot.execute("""
                    SELECT DISTINCT r.round_number FROM round_results r
                    JOIN demand_commitments d ON d.round_number=r.round_number
                    WHERE d.completed_at IS NULL
                """).fetchall()
                if unfinished:
                    raise ValueError("Round results are still being finalized. Prepare downloads after round processing finishes.")
                metadata = {
                    "format": BACKUP_FORMAT, "format_version": BACKUP_VERSION,
                    "classroom_id": classroom_id, "classroom_name": name,
                    "captured_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    "database_file": "classroom.sqlite3", "table_counts": _counts(snapshot),
                    "latest_completed_round": snapshot.execute("SELECT COALESCE(MAX(round_number),0) FROM round_results").fetchone()[0],
                    "schema_sha256": sha256("\n".join(row[0] for row in snapshot.execute(
                        "SELECT sql FROM sqlite_master WHERE sql IS NOT NULL ORDER BY type,name")).encode()).hexdigest(),
                    "contains_password_hashes": True,
                }
            database_bytes = path.read_bytes()
            metadata["database_sha256"] = sha256(database_bytes).hexdigest()
            metadata["database_bytes"] = len(database_bytes)
    # Release the live read transaction before workbook/ZIP generation.
    with closing(sqlite3.connect(":memory:")) as snapshot:
        snapshot.deserialize(database_bytes)
        snapshot.execute("PRAGMA query_only=ON")
        try:
            workbook = _workbook_bytes(snapshot, metadata)
            excel_error = None
        except ValueError as error:
            workbook = None
            excel_error = str(error)
    archive = BytesIO()
    with ZipFile(archive, "w", compression=ZIP_DEFLATED) as zipped:
        zipped.writestr("classroom.sqlite3", database_bytes)
        zipped.writestr("manifest.json", json.dumps(metadata, indent=2, ensure_ascii=False))
        zipped.writestr("RESTORE_README.txt", RESTORE_README)
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_")[:60] or "classroom"
    stamp = metadata["captured_at_utc"].replace("-", "").replace(":", "").replace("+0000", "Z")
    prefix = f"kiki_{slug}_{classroom_id[:8]}_{stamp}"
    return {"metadata": metadata, "excel": workbook, "excel_error": excel_error, "backup": archive.getvalue(),
            "excel_name": prefix + "_data.xlsx", "backup_name": prefix + "_backup.zip"}


def restore_backup_to_new_file(archive_bytes, destination):
    """Validate and recover only to a new, isolated path. Never overwrite a file."""
    destination = Path(destination)
    if destination.exists() or destination.is_symlink():
        raise FileExistsError("Recovery requires a new file. Existing files are never overwritten.")
    if len(archive_bytes) > MAX_BACKUP_BYTES:
        raise ValueError("Backup exceeds the supported recovery size.")
    with ZipFile(BytesIO(archive_bytes)) as zipped:
        names = zipped.namelist()
        expected = {"classroom.sqlite3", "manifest.json", "RESTORE_README.txt"}
        if len(names) != len(expected) or set(names) != expected:
            raise ValueError("Invalid backup contents.")
        if zipped.getinfo("manifest.json").file_size > 1024 * 1024:
            raise ValueError("Invalid backup manifest size.")
        if zipped.getinfo("classroom.sqlite3").file_size > MAX_BACKUP_BYTES:
            raise ValueError("Database exceeds the supported recovery size.")
        metadata = json.loads(zipped.read("manifest.json"))
        if (metadata.get("format") != BACKUP_FORMAT or metadata.get("format_version") != BACKUP_VERSION
                or metadata.get("database_file") != "classroom.sqlite3"):
            raise ValueError("Unsupported backup format.")
        data = zipped.read("classroom.sqlite3")
    if len(data) != metadata.get("database_bytes") or sha256(data).hexdigest() != metadata.get("database_sha256"):
        raise ValueError("Backup checksum mismatch. Recovery was not performed.")
    with closing(sqlite3.connect(":memory:")) as connection:
        connection.deserialize(data)
        connection.execute("PRAGMA trusted_schema=OFF")
        connection.execute("PRAGMA query_only=ON")
        if connection.execute("PRAGMA integrity_check").fetchall() != [("ok",)]:
            raise ValueError("Backup database integrity check failed.")
        if _counts(connection) != metadata.get("table_counts"):
            raise ValueError("Backup table counts do not match its manifest.")
        if not connection.execute("SELECT 1 FROM users WHERE role='admin' AND is_active=1 LIMIT 1").fetchone():
            raise ValueError("Backup has no active administrator.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation also prevents a concurrent request from replacing a file.
    with destination.open("xb") as output:
        output.write(data)
        output.flush()
    return metadata
