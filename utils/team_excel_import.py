"""Create-only classroom roster imports. Never update existing teams or logins."""

from dataclasses import asdict, fields
from io import BytesIO
import json
import zipfile

from openpyxl import load_workbook

from data.classroom_start import COMMON_PROFILE, initial_state
from models.schemas import ProductLine, PROJECT_SLOT_NAMES, build_product_id
from utils.database import get_connection
from utils.repository import _empty_development_project_slot, load_market_report
from utils.security import hash_password, validate_password_strength

HEADERS = ("team_name", "username", "password")
MAX_TEAMS = 200


def read_team_workbook(content: bytes) -> list[dict]:
    if len(content) > 2_000_000:
        raise ValueError("Workbook must be smaller than 2 MB.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20_000_000:
                raise ValueError("Workbook is too large when unpacked.")
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise ValueError("Cannot read this file. Upload a valid .xlsx workbook.") from error
    try:
        if "Teams" not in workbook.sheetnames:
            raise ValueError("The workbook must contain a sheet named Teams.")
        sheet = workbook["Teams"]
        if (sheet.max_row or 0) > MAX_TEAMS + 1 or (sheet.max_column or 0) > 3:
            raise ValueError("Teams must have only three columns and at most 200 team rows.")
        cells = sheet.iter_rows(min_row=1, max_col=4, max_row=MAX_TEAMS + 2)
        header = next(cells, ())
        if tuple(cell.value for cell in header[:3]) != HEADERS or header[3].value is not None:
            raise ValueError("Row 1 must be: team_name, username, password (in this order).")
        rows = []
        for number, row in enumerate(cells, start=2):
            values = [cell.value for cell in row]
            if all(value in (None, "") for value in values):
                continue
            if number > MAX_TEAMS + 1 or values[3] not in (None, ""):
                raise ValueError("Teams must have only three columns and at most 200 team rows.")
            row = row[:3]
            values = [cell.value if cell.value is not None else "" for cell in row]
            if any(cell.data_type == "f" for cell in row):
                raise ValueError(f"Row {number}: formulas are not allowed.")
            if any(not isinstance(value, str) for value in values):
                raise ValueError(f"Row {number}: fill all three cells as Text, not numbers or dates.")
            rows.append(dict(zip(HEADERS, values), row=number))
        if not rows:
            raise ValueError("Enter at least one team in the Teams sheet.")
        return rows
    finally:
        workbook.close()


def _existing_names(connection):
    usernames = {row[0].casefold() for row in connection.execute("SELECT username FROM users")}
    teams = set()
    # Include historical and inactive teams, even if their account was removed.
    for table in ("users", "team_states", "product_lines", "team_decisions",
                  "product_decisions", "product_forecasts", "product_development_projects",
                  "round_results", "product_round_results", "forecast_accuracy_results"):
        teams.update(row[0].casefold() for row in connection.execute(
            f"SELECT DISTINCT team_name FROM {table} WHERE team_name IS NOT NULL"))
    return usernames, teams


def validate_team_rows(rows: list[dict], connection=None, *, allow_classroom_passwords=False) -> list[dict]:
    if connection is None:
        with get_connection() as connection:
            return validate_team_rows(rows, connection, allow_classroom_passwords=allow_classroom_passwords)
    usernames, teams = _existing_names(connection)
    product_ids = {row[0] for row in connection.execute("SELECT product_id FROM product_lines")}
    team_keys = {build_product_id(name, "A") for name in teams}
    seen_users, seen_teams, seen_keys = set(), set(), set()
    errors = []
    if not 1 <= len(rows) <= MAX_TEAMS:
        return [{"row": 0, "error": "Import between 1 and 200 teams."}]
    for row in rows:
        team, username, password = (row[key] for key in HEADERS)
        messages = []
        for key in HEADERS:
            value = row[key]
            if not value or value != value.strip() or any(ord(c) < 32 for c in value):
                messages.append(f"{key}: required; no surrounding spaces or control characters.")
            if len(value) > 128:
                messages.append(f"{key}: maximum 128 characters.")
        # An instructor may opt in to simple student passwords for this import only.
        if not allow_classroom_passwords:
            messages.extend(validate_password_strength(password))
        user_key, team_key = username.casefold(), team.casefold()
        product_key = build_product_id(team, "A")
        if user_key in usernames or user_key in seen_users:
            messages.append("Username already exists or is repeated in this workbook.")
        if team_key in teams or team_key in seen_teams:
            messages.append("Team name already exists or is repeated in this workbook.")
        if (product_key in product_ids or product_key in team_keys or product_key in seen_keys):
            messages.append("Team name conflicts with another team's internal product identifiers.")
        if not any(c.isalnum() for c in team):
            messages.append("Team name must contain a letter or number.")
        seen_users.add(user_key)
        seen_teams.add(team_key)
        seen_keys.add(product_key)
        if messages:
            errors.append({"row": row["row"], "error": " ".join(messages)})
    return errors


def _insert_new(connection, table: str, payload: dict) -> None:
    # Table/column names come exclusively from our own schema objects, never Excel.
    columns = ",".join(payload)
    placeholders = ",".join("?" for _ in payload)
    connection.execute(f"INSERT INTO {table} ({columns}) VALUES ({placeholders})", tuple(payload.values()))


def create_teams_from_workbook(rows: list[dict], *, allow_classroom_passwords=False) -> int:
    current_round = load_market_report().round_number
    hashes = [hash_password(row["password"]) for row in rows]
    with get_connection() as connection:
        # Recheck under the write lock: concurrent imports cannot overwrite a team.
        connection.execute("BEGIN IMMEDIATE")
        errors = validate_team_rows(rows, connection, allow_classroom_passwords=allow_classroom_passwords)
        if errors:
            raise ValueError("Import canceled. Refresh the preview and resolve the reported conflicts.")
        for row, password_hash in zip(rows, hashes):
            team = row["team_name"]
            _insert_new(connection, "users", dict(username=row["username"], password_hash=password_hash,
                        role="team_leader", team_name=team, is_active=1))
            state = asdict(initial_state(team))
            for key in ("completed_rounds", "last_decision", "open_material_orders"):
                state[key + "_json"] = json.dumps(state.pop(key))
            _insert_new(connection, "team_states", state)
            product_fields = {field.name for field in fields(ProductLine)}
            for template in COMMON_PROFILE.suggested_product_templates:
                payload = {key: value for key, value in asdict(template).items() if key in product_fields}
                product = ProductLine(product_id=build_product_id(team, template.slot_name),
                                      team_name=team, **payload)
                _insert_new(connection, "product_lines", asdict(product))
            for slot in PROJECT_SLOT_NAMES:
                project = _empty_development_project_slot(team, slot, current_round)
                _insert_new(connection, "product_development_projects", project.to_dict())
    return len(rows)
