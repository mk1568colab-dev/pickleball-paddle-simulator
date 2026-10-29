"""Read-only exports, classroom isolation and new-file recovery regressions."""

from contextlib import contextmanager, closing
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TEMP = tempfile.TemporaryDirectory()
BASE = Path(TEMP.name) / "demo.db"
os.environ["SIMULATOR_DB_PATH"] = str(BASE)
os.environ["SIMULATOR_ENV"] = "dev"
os.environ["SIMULATOR_ENABLE_DEMO_ACCOUNTS"] = "false"

from openpyxl import load_workbook
from scripts.export_demo_data import seed_export_demo
from utils.game_exports import prepare_classroom_downloads, restore_backup_to_new_file
from utils.classrooms import AUTH_CLASSROOM_KEY, database_scope
from utils.security import verify_password


@contextmanager
def session_for(classroom, username="export_admin"):
    state = {"authenticated_username": username, AUTH_CLASSROOM_KEY: classroom["id"]}
    with patch("utils.classrooms._session", return_value=state), patch("utils.game_exports._session", return_value=state):
        yield


def database_dump(path):
    with closing(sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro", uri=True)) as connection:
        return "\n".join(connection.iterdump())


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.classes = seed_export_demo(BASE)
        cls.bundles = []
        cls.before = [database_dump(c["path"]) for c in cls.classes]
        cls.hashes = [sha256(c["path"].read_bytes()).hexdigest() for c in cls.classes]
        for classroom in cls.classes:
            with session_for(classroom):
                cls.bundles.append(prepare_classroom_downloads())

    def test_read_only_accounts_and_progress(self):
        self.assertEqual([database_dump(c["path"]) for c in self.classes], self.before)
        self.assertEqual([sha256(c["path"].read_bytes()).hexdigest() for c in self.classes], self.hashes)

    def test_every_saved_table_and_round_in_excel(self):
        for classroom, bundle in zip(self.classes, self.bundles):
            workbook = load_workbook(BytesIO(bundle["excel"]))
            self.assertEqual(workbook["Team Results"].max_row - 1, 2 * classroom["rounds"])
            self.assertEqual(workbook["Firm Decisions"].max_row - 1, 2 * classroom["rounds"])
            self.assertEqual({row[0].value for row in workbook["Team Results"].iter_rows(min_row=2)},
                             set(range(1, classroom["rounds"] + 1)))
            for sheet, table, count, scope in workbook["Table Index"].iter_rows(min_row=2, values_only=True):
                if table in bundle["metadata"]["table_counts"]:
                    self.assertEqual(count, bundle["metadata"]["table_counts"][table])
                    self.assertEqual(workbook[sheet].max_row - 1, count)
            self.assertTrue(workbook["Team Results"].auto_filter.ref)
            self.assertEqual(workbook["Team Results"].freeze_panes, "C2")

    def test_classroom_isolation_and_filenames(self):
        for classroom, bundle in zip(self.classes, self.bundles):
            workbook = load_workbook(BytesIO(bundle["excel"]))
            names = {row[1].value for row in workbook["Team Results"].iter_rows(min_row=2)}
            self.assertEqual(names, {f"Demo {classroom['section']} Team 1", f"Demo {classroom['section']} Team 2"})
            self.assertIn(classroom["section"], bundle["excel_name"])
            self.assertEqual(bundle["metadata"]["classroom_id"], classroom["id"])
            self.assertNotIn(str(BASE.parent), json.dumps(bundle["metadata"]))

    def test_excel_excludes_credentials_and_formulas(self):
        for bundle in self.bundles:
            with ZipFile(BytesIO(bundle["excel"])) as zipped:
                xml = b"".join(zipped.read(n) for n in zipped.namelist() if n.endswith(".xml"))
            self.assertNotIn(b"password_hash", xml)
            self.assertNotIn(b"PreviewOnly123", xml)
            workbook = load_workbook(BytesIO(bundle["excel"]))
            self.assertFalse(any(c.data_type == 'f' for s in workbook for r in s for c in r))

    def test_backup_restore_exact_state_and_passwords(self):
        for classroom, bundle, before in zip(self.classes, self.bundles, self.before):
            target = Path(TEMP.name) / f"recovered_{classroom['section']}.db"
            meta = restore_backup_to_new_file(bundle["backup"], target)
            self.assertEqual(database_dump(target), before)
            self.assertEqual(meta['latest_completed_round'], classroom['rounds'])
            with closing(sqlite3.connect(target)) as connection:
                hash_value = connection.execute("SELECT password_hash FROM users WHERE username='export_admin'").fetchone()[0]
                self.assertTrue(verify_password("PreviewOnly123!", hash_value))
                self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], 'ok')
            with self.assertRaises(FileExistsError):
                restore_backup_to_new_file(bundle["backup"], target)
            self.assertEqual(database_dump(target), before)

    def test_anonymous_student_inactive_and_wrong_class_denied(self):
        with patch("utils.game_exports._session", return_value=None), self.assertRaises(PermissionError):
            prepare_classroom_downloads()
        with session_for(self.classes[0], "demo_022_1"), self.assertRaises(PermissionError):
            prepare_classroom_downloads()
        with session_for(self.classes[1], "demo_022_1"), self.assertRaises(PermissionError):
            prepare_classroom_downloads()
        copy = Path(TEMP.name) / "inactive.db"
        restore_backup_to_new_file(self.bundles[0]['backup'], copy)
        with closing(sqlite3.connect(copy)) as connection:
            connection.execute("UPDATE users SET is_active=0 WHERE username='export_admin'")
            connection.commit()
        with session_for(self.classes[0]), database_scope(copy), self.assertRaises(PermissionError):
            prepare_classroom_downloads()

    def test_wal_snapshot_includes_commits_without_checkpoint(self):
        copy = Path(TEMP.name) / "wal.db"
        restore_backup_to_new_file(self.bundles[0]['backup'], copy)
        with closing(sqlite3.connect(copy)) as writer:
            writer.execute("PRAGMA journal_mode=WAL")
            writer.execute("PRAGMA wal_autocheckpoint=0")
            writer.execute("UPDATE users SET username='wal_only_student' WHERE username='demo_022_1'")
            writer.commit()
            self.assertTrue(Path(str(copy) + "-wal").exists())
            with session_for(self.classes[0]), database_scope(copy):
                bundle = prepare_classroom_downloads()
            target = Path(TEMP.name) / "wal_recovered.db"
            restore_backup_to_new_file(bundle['backup'], target)
            self.assertEqual(database_dump(target), database_dump(copy))

    def test_partial_round_is_rejected_without_writing(self):
        copy = Path(TEMP.name) / "partial_round.db"
        restore_backup_to_new_file(self.bundles[0]['backup'], copy)
        with closing(sqlite3.connect(copy)) as connection:
            connection.execute("UPDATE demand_commitments SET completed_at=NULL WHERE round_number=3")
            connection.commit()
        before = database_dump(copy)
        with session_for(self.classes[0]), database_scope(copy), self.assertRaisesRegex(ValueError, 'finalized'):
            prepare_classroom_downloads()
        self.assertEqual(database_dump(copy), before)

    def test_backup_remains_available_when_excel_is_too_large(self):
        with session_for(self.classes[0]), patch('utils.game_exports._workbook_bytes', side_effect=ValueError('Excel limits')):
            bundle = prepare_classroom_downloads()
        self.assertIsNone(bundle['excel'])
        self.assertEqual(bundle['excel_error'], 'Excel limits')
        target = Path(TEMP.name) / 'backup_only.db'
        restore_backup_to_new_file(bundle['backup'], target)
        self.assertEqual(database_dump(target), self.before[0])

    def test_formula_injection_and_long_text_preserved(self):
        copy = Path(TEMP.name) / "text.db"
        restore_backup_to_new_file(self.bundles[0]['backup'], copy)
        long_value = "A" * 40000
        with closing(sqlite3.connect(copy)) as connection:
            connection.execute("UPDATE classroom_rounds SET notes=? WHERE round_number=1", ('=HYPERLINK("https://example.invalid","click")',))
            connection.execute("UPDATE classroom_rounds SET notes=? WHERE round_number=2", (long_value,))
            connection.commit()
        with session_for(self.classes[0]), database_scope(copy):
            bundle = prepare_classroom_downloads()
        workbook = load_workbook(BytesIO(bundle['excel']))
        formula = workbook['Round Status'].cell(2, 3)
        self.assertTrue(formula.value.startswith('=HYPERLINK'))
        self.assertEqual(formula.data_type, 's')
        self.assertEqual(''.join(r[4] for r in workbook['Text Chunks'].iter_rows(min_row=2, values_only=True)), long_value)

    def test_tampering_unsupported_archive_and_overwrite_refused(self):
        bundle = self.bundles[0]
        with self.assertRaises(FileExistsError):
            restore_backup_to_new_file(bundle['backup'], BASE)
        with ZipFile(BytesIO(bundle['backup'])) as source:
            original = {n: source.read(n) for n in source.namelist()}
        for case in ('checksum', 'extra_path', 'count', 'version', 'duplicate'):
            files = dict(original)
            meta = json.loads(files['manifest.json'])
            if case == 'checksum': files['classroom.sqlite3'] += b'bad'
            if case == 'extra_path': files['../../unexpected.txt'] = b'bad'
            if case == 'count': meta['table_counts']['users'] = 999
            if case == 'version': meta['format_version'] = 999
            files['manifest.json'] = json.dumps(meta).encode()
            out = BytesIO()
            with ZipFile(out, 'w', ZIP_DEFLATED) as zipped:
                for name, data in files.items(): zipped.writestr(name, data)
                if case == 'duplicate': zipped.writestr('manifest.json', files['manifest.json'])
            target = Path(TEMP.name) / f"bad_{case}.db"
            with self.assertRaises(ValueError): restore_backup_to_new_file(out.getvalue(), target)
            self.assertFalse(target.exists())

    @patch('streamlit.page_link')
    def test_admin_page_and_student_direct_access(self, page_link):
        from streamlit.testing.v1 import AppTest
        page = ROOT / 'pages/14_Data_Export_Backup.py'
        for classroom in self.classes:
            app = AppTest.from_file(str(page), default_timeout=30)
            app.session_state['authenticated_username'] = 'export_admin'
            app.session_state[AUTH_CLASSROOM_KEY] = classroom['id']
            app.run()
            self.assertFalse(app.exception)
            next(b for b in app.button if b.label == 'Prepare Downloads').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.get('download_button')), 2)
            bundle = app.session_state['classroom_download_bundle']['bundle']
            self.assertEqual(bundle['metadata']['classroom_id'], classroom['id'])
        app = AppTest.from_file(str(page), default_timeout=30)
        app.session_state['authenticated_username'] = 'demo_022_1'
        app.session_state[AUTH_CLASSROOM_KEY] = 'default'
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('download_button')), 0)
        self.assertTrue(any('admin users only' in e.value for e in app.error))


if __name__ == '__main__':
    unittest.main(verbosity=2)
