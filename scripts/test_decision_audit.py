"""Isolated regression tests for durable decision revisions and export coverage."""

from concurrent.futures import ThreadPoolExecutor
from contextlib import closing, contextmanager
from dataclasses import replace
from datetime import datetime
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TEMP = tempfile.TemporaryDirectory()
os.environ['SIMULATOR_DB_PATH'] = str(Path(TEMP.name) / 'unused.db')
os.environ['SIMULATOR_ENV'] = 'dev'
os.environ['SIMULATOR_ENABLE_DEMO_ACCOUNTS'] = 'false'

from data.classroom_start import COMMON_PROFILE_NAME
from models.schemas import ProductDecision, ProductDevelopmentProject, TeamDecision
from utils import repository as repo
from utils.classrooms import AUTH_CLASSROOM_KEY, database_scope
from utils.database import get_connection, initialize_database
from utils.decision_audit import load_audit_history, save_audited_submission
from utils.game_exports import prepare_classroom_downloads, restore_backup_to_new_file
from utils.submission_fallback import complete_round_inputs, save_fallback_decisions


@contextmanager
def session(username):
    state = {'authenticated_username': username, AUTH_CLASSROOM_KEY: 'default'} if username else {}
    with patch('utils.classrooms._session', return_value=state), patch('utils.decision_audit._session', return_value=state), \
         patch('utils.game_exports._session', return_value=state):
        yield


def records(connection, exclude_audit=False):
    names = [r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
    return {name: [tuple(r) for r in connection.execute(f'SELECT * FROM "{name}" ORDER BY rowid')]
            for name in names if not (exclude_audit and name.startswith('decision_audit_'))}


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.path = Path(TEMP.name) / (self.id().split('.')[-1] + '.db')
        scope = database_scope(self.path)
        scope.__enter__()
        self.addCleanup(scope.__exit__, None, None, None)
        self.admin = repo.create_user('audit_admin', 'test-only-hash', 'admin', None, True)
        self.student = repo.create_user('audit_student', 'test-only-hash', 'team_leader', 'Audit Team', True)
        repo.create_user('other_student', 'test-only-hash', 'team_leader', 'Other Team', True)
        self.lines = repo.load_product_lines('Audit Team')
        self.firm = TeamDecision.from_dict(dict(team_name='Audit Team', archetype=COMMON_PROFILE_NAME,
            raw_material_order_qty=240, supplier_mix_offshore_pct=30, supplier_mix_balanced_pct=50, supplier_mix_premium_pct=20))
        self.products = [ProductDecision.from_dict(dict(team_name='Audit Team', slot_name=p.slot_name,
            product_id=p.product_id, product_name=p.product_name, target_segment=p.target_segment,
            is_active=p.is_active, selling_price_per_unit=150 if p.is_active else 0,
            planned_production_units=110 if p.is_active else 0, forecast_units=120 if p.is_active else 0,
            qc_budget_per_unit=3)) for p in self.lines]
        self.projects = [ProductDevelopmentProject.from_dict(dict(team_name='Audit Team', project_slot_name='P1',
            project_name='Next Paddle', target_segment='premium', investment_this_round=500,
            testing_intensity=.4, planned_launch_round=4))]

    def submit(self, price=150, round_number=1, admin=False, investment=500):
        products = [replace(p, selling_price_per_unit=price) if p.slot_name == 'A' else p for p in self.products]
        projects = [replace(p, investment_this_round=investment) for p in self.projects]
        actor = self.admin if admin else self.student
        with session(actor.username):
            save_audited_submission(self.firm, products, projects, self.lines, round_number, actor.user_id)

    def history(self):
        with session('audit_admin'):
            return load_audit_history()

    def test_all_revisions_actor_values_and_timestamp(self):
        for price in (150, 180, 170): self.submit(price)
        context, events, changes = self.history()
        self.assertEqual([e['revision'] for e in events], [3, 2, 1])
        self.assertEqual([json.loads(e['after_json'])['products']['A']['selling_price_per_unit'] for e in events], [170, 180, 150])
        last = events[0]
        self.assertEqual(last['actor_username'], 'audit_student')
        self.assertEqual(last['actor_user_id'], self.student.user_id)
        self.assertEqual(last['source'], 'student_submission')
        self.assertEqual(last['comparison_basis'], 'same_round_saved_plan')
        self.assertIsNotNone(datetime.fromisoformat(last['occurred_at_utc']).utcoffset())
        delta = next(c for c in changes if c['event_id'] == last['event_id'] and c['field_name'] == 'selling_price_per_unit')
        self.assertEqual((json.loads(delta['before_value_json']), json.loads(delta['after_value_json'])), (180, 170))
        self.assertEqual(repo.load_product_decisions(1, 'Audit Team')[0].selling_price_per_unit, 170)

    def test_unchanged_resubmission_is_an_event(self):
        self.submit(); self.submit()
        _, events, changes = self.history()
        self.assertEqual(len(events), 2)
        self.assertEqual(events[0]['change_count'], 0)
        self.assertFalse(any(c['event_id'] == events[0]['event_id'] for c in changes))

    def test_pipeline_controls_are_historical(self):
        self.submit(investment=500); self.submit(investment=900)
        _, events, changes = self.history()
        delta = next(c for c in changes if c['event_id'] == events[0]['event_id'] and c['field_name'] == 'investment_this_round')
        self.assertEqual(delta['decision_group'], 'Development')
        self.assertEqual(json.loads(delta['before_value_json']), 500)
        self.assertEqual(json.loads(delta['after_value_json']), 900)
        self.assertEqual(json.loads(events[1]['after_json'])['pipeline']['P1']['investment_this_round'], 500)

    def test_administrator_is_not_impersonated_as_student(self):
        self.submit(admin=True)
        event = self.history()[1][0]
        self.assertEqual((event['source'], event['actor_username']), ('admin_submission', 'audit_admin'))

    def test_automatic_carry_records_system_admin_and_zero_development(self):
        self.submit()
        inputs = complete_round_inputs(2, [], [], repo.load_product_development_projects(), repo.load_product_lines(), repo.load_team_states())
        with session('audit_admin'):
            save_fallback_decisions(2, inputs[0], inputs[1], inputs[4], inputs[2])
        _, events, changes = self.history()
        event = next(e for e in events if e['team_name'] == 'Audit Team' and e['round_number'] == 2)
        self.assertEqual((event['source'], event['actor_username'], event['actor_user_id']), ('automatic_carry_forward', 'System', None))
        self.assertEqual(event['initiated_by_username'], 'audit_admin')
        self.assertEqual(event['prior_round_number'], 1)
        after = json.loads(event['after_json'])['pipeline']['P1']
        self.assertEqual([after[k] for k in ('investment_this_round', 'testing_intensity', 'launch_now', 'cancel_now')], [0, 0, False, False])
        with get_connection() as connection:
            self.assertIsNone(connection.execute("SELECT submitted_by_user_id FROM team_decisions WHERE team_name='Audit Team' AND round_number=2").fetchone()[0])

    def test_failed_submission_and_failed_audit_both_roll_back(self):
        self.submit()
        with get_connection() as connection: before = records(connection)
        with patch('utils.repository.save_product_lines', side_effect=RuntimeError('simulated failure')), self.assertRaises(RuntimeError):
            self.submit(199)
        with get_connection() as connection: self.assertEqual(records(connection), before)
        with get_connection() as connection:
            connection.execute("CREATE TRIGGER reject_audit BEFORE INSERT ON decision_audit_events BEGIN SELECT RAISE(ABORT,'test audit failure'); END")
        with self.assertRaises(sqlite3.IntegrityError): self.submit(199)
        with get_connection() as connection: self.assertEqual(records(connection), before)

    def test_wrong_actor_wrong_team_and_student_automation_denied(self):
        with session('other_student'), self.assertRaises(PermissionError):
            save_audited_submission(self.firm, self.products, [], self.lines, 1, self.student.user_id)
        with session('audit_student'), self.assertRaises(PermissionError):
            save_audited_submission(replace(self.firm, team_name='Other Team'), [], [], [], 1, self.student.user_id)
        with session(None), self.assertRaises(PermissionError):
            save_audited_submission(self.firm, self.products, [], self.lines, 1, self.student.user_id)
        with session('audit_student'), self.assertRaises(PermissionError):
            save_fallback_decisions(1, [self.firm], self.products, ['Audit Team'])
        self.assertEqual(self.history()[1], [])
        self.assertIsNone(repo.load_team_decision(1, 'Audit Team'))

    def test_migration_preserves_existing_records_and_does_not_backfill(self):
        repo.save_team_decision(self.firm, 1, self.student.user_id)
        repo.save_product_decisions(self.products, 1)
        with get_connection() as connection:
            for table in ('decision_audit_changes', 'decision_audit_events', 'decision_audit_context'):
                connection.execute(f'DROP TABLE {table}')
            before = records(connection)
        initialize_database()
        with get_connection() as connection:
            self.assertEqual(records(connection, exclude_audit=True), before)
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM decision_audit_events').fetchone()[0], 0)
        self.submit(170)
        event = self.history()[1][0]
        self.assertEqual(event['revision'], 1)
        self.assertEqual(json.loads(event['before_json'])['products']['A']['selling_price_per_unit'], 150)

    def test_append_only_and_reset_retains_history_with_new_game_id(self):
        self.submit()
        old = self.history()
        for table in ('decision_audit_events', 'decision_audit_changes'):
            for sql in (f'DELETE FROM {table}', f'UPDATE {table} SET team_name=team_name'):
                with self.assertRaises(sqlite3.IntegrityError):
                    with get_connection() as connection: connection.execute(sql)
        repo.factory_reset_game_data()
        new = self.history()
        self.assertEqual(new[1:], old[1:])
        self.assertNotEqual(new[0]['game_id'], old[0]['game_id'])
        self.submit(190)
        events = self.history()[1]
        self.assertEqual([e['revision'] for e in events], [1, 1])
        self.assertNotEqual(events[0]['game_id'], events[1]['game_id'])

    def test_excel_and_recovery_include_every_audit_event_and_change(self):
        from openpyxl import load_workbook
        self.submit(); self.submit(180); self.submit(170)
        with session('audit_admin'): bundle = prepare_classroom_downloads()
        workbook = load_workbook(BytesIO(bundle['excel']))
        self.assertEqual(workbook['Decision Audit Events'].max_row - 1, 3)
        self.assertEqual(workbook['Decision Audit Changes'].max_row - 1, len(self.history()[2]))
        restored = Path(TEMP.name) / 'audit_restored.db'
        restore_backup_to_new_file(bundle['backup'], restored)
        with closing(sqlite3.connect(restored)) as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM decision_audit_events').fetchone()[0], 3)
            with self.assertRaises(sqlite3.IntegrityError): connection.execute('DELETE FROM decision_audit_events')

    def test_two_concurrent_submissions_have_a_complete_ordered_chain(self):
        self.submit()
        def save(price):
            with database_scope(self.path):
                products = [replace(p, selling_price_per_unit=price) if p.slot_name == 'A' else p for p in self.products]
                save_audited_submission(self.firm, products, self.projects, self.lines, 1, self.student.user_id)
        with session('audit_student'), ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(save, [180, 170]))
        events = list(reversed(self.history()[1]))
        self.assertEqual([e['revision'] for e in events], [1, 2, 3])
        self.assertEqual(events[0]['after_json'], events[1]['before_json'])
        self.assertEqual(events[1]['after_json'], events[2]['before_json'])

    @patch('streamlit.page_link')
    def test_admin_page_filters_and_student_denial(self, page_link):
        from streamlit.testing.v1 import AppTest
        self.submit(); self.submit(180)
        # AppTest runs in another context; point its base only at this disposable DB.
        with patch.dict(os.environ, {'SIMULATOR_DB_PATH': str(self.path)}):
            app = AppTest.from_file(str(ROOT / 'pages/15_Decision_Audit_Log.py'), default_timeout=30)
            app.session_state['authenticated_username'] = 'audit_admin'
            app.session_state[AUTH_CLASSROOM_KEY] = 'default'
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.dataframe[0].value), 2)
            next(s for s in app.selectbox if s.label == 'Source').select('Automatic carry-forward').run()
            self.assertTrue(any('No revisions' in i.value for i in app.info))
            denied = AppTest.from_file(str(ROOT / 'pages/15_Decision_Audit_Log.py'), default_timeout=30)
            denied.session_state['authenticated_username'] = 'audit_student'
            denied.run()
            self.assertFalse(denied.exception)
            self.assertEqual(len(denied.dataframe), 0)
            self.assertTrue(denied.error)


if __name__ == '__main__':
    unittest.main(verbosity=2)
